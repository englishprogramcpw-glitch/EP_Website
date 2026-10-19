import sys
import os
import urllib.request
import urllib.parse
import http.cookiejar
import json
import re
import subprocess

sys.stdout.reconfigure(encoding='utf-8')

# Paths
base_dir = os.path.dirname(os.path.abspath(__file__))
db_path = os.path.join(base_dir, 'data', 'db_staff.json')
img_dir = os.path.join(base_dir, 'img', 'teachers')
env_path = os.path.join(base_dir, '..', 'EP_StudentSync', '.env')

ADMIN_ROLE_EN_MAP = {
    'ที่ปรึกษาโครงการฯ': 'Advisor of English Program',
    'ที่ปรึกษาโครงการฯ ภาคภาษาอังกฤษ': 'Advisor of English Program',
    'หัวหน้าโครงการฯ': 'Head of English Program',
    'หัวหน้าโครงการฯ ภาคภาษาอังกฤษ': 'Head of English Program',
    'ผู้ช่วยหัวหน้าโครงการฯ': 'Assistant Head of English Program',
    'หัวหน้างานวิชาการ': 'Head of Academic Affairs',
    'หัวหน้างานปกครอง': 'Head of Student Affairs',
    'หัวหน้างานปกครอง โครงการฯ': 'Head of Student Affairs',
    'หัวหน้างานสำนักงาน และเจ้าหน้าที่งานการเงิน': 'Head of Office & Financial Officer',
    'หัวหน้างานสำนักงาน, เจ้าหน้าที่งานการเงิน': 'Head of Office & Financial Officer',
    'หัวหน้างานสำนักงานและเจ้าหน้าที่งานการเงิน': 'Head of Office & Financial Officer',
    'เจ้าหน้าที่ฝ่ายวิชาการ': 'Academic Officer',
    'เจ้าหน้าที่งานวิชาการ': 'Academic Officer',
    'เจ้าหน้าที่จัดเตรียมเอกสาร': 'Documentation Officer',
    'เจ้าหน้าที่งานเอกสาร': 'Documentation Officer',
    'เจ้าหน้าที่บรรณารักษ์': 'Librarian',
    'เจ้าหน้าที่งานธุรการ': 'Registrar / Administrative Officer',
    'เจ้าหน้าที่งานพัสดุ และอาคารสถานที่': 'Procurement & Facilities Officer',
    'เจ้าหน้าที่งานพัสดุและอาคารสถานที่': 'Procurement & Facilities Officer',
    'เจ้าหน้าที่งานสารสนเทศ และประชาสัมพันธ์': 'IT & Public Relations Officer',
    'เจ้าหน้าที่งานสารสนเทศและประชาสัมพันธ์': 'IT & Public Relations Officer',
    'ผู้ประสานงานวิชาการระดับมัธยมศึกษา': 'Secondary Academic Coordinator',
    'ผู้ประสานงานวิชาการมัธยมฯ': 'Secondary Academic Coordinator',
    'เจ้าหน้าที่ผู้ประสานงานวิชาการ ระดับมัธยมศึกษา': 'Secondary Academic Coordinator',
    'เจ้าหน้าที่งานห้องสมุด': 'Library Officer',
}

SUBJECT_MAP_EN = {
    'คณิตศาสตร์': 'Mathematics',
    'วิทยาศาสตร์': 'Science',
    'สังคมศึกษา': 'Social Studies',
    'สังคมศึกษาศาสนาและวัฒนธรรม': 'Social Studies',
    'ประวัติศาสตร์': 'History',
    'ภาษาไทย': 'Thai Language',
    'ภาษาอังกฤษ': 'English',
    'ศิลปะ': 'Arts',
    'ศิลปะและหัตถกรรม': 'Arts & Crafts',
    'การงานอาชีพ': 'Occupational Works',
    'สุขศึกษา': 'Health Education',
    'สุขศึกษาและพลศึกษา': 'Health & Physical Education',
    'พลศึกษา': 'Physical Education',
    'ดนตรี': 'Music',
    'นาฏศิลป์': 'Thai Dance',
    'คอมพิวเตอร์': 'Computer Science',
    'เทคโนโลยี': 'Technology',
    'ฟิสิกส์': 'Physics',
    'เคมี': 'Chemistry',
    'ชีววิทยา': 'Biology',
    'วิทยาการคำนวณ': 'Computing Science',
    'ภาษาจีน': 'Chinese',
    'แนะแนว': 'Guidance Counseling',
    'กิจกรรมพัฒนาผู้เรียน': 'Student Development Activities',
}

def clean_subject_name(s):
    if not s:
        return ""
    s = s.replace('\n', ' ').replace('\r', ' ').replace('\t', ' ').strip()
    s = " ".join(s.split())
    return s

def split_subjects(subj_str):
    if not subj_str:
        return []
    s = subj_str.replace(' และ ', ',').replace(' และ', ',').replace('และ', ',')
    raw_parts = [p.strip() for p in re.split(r'[,，\n]', s) if p.strip()]
    
    parts = []
    for p in raw_parts:
        clean_p = re.sub(r'[ปม]\.\s*[\d\s\-\,\/]+', '', p).strip(" ,")
        if clean_p:
            parts.append(p)
        else:
            if parts:
                parts[-1] = parts[-1] + ", " + p
            else:
                parts.append(p)
    return parts

def format_teaching_role(teach, assist):
    main_parts = split_subjects(teach)
    assist_parts = split_subjects(assist)
    
    lines = []
    if main_parts:
        first_part = clean_subject_name(main_parts[0])
        if first_part in ["ครูผู้สอน", "ครู"] or not first_part:
            lines.append("ครูผู้สอน")
        else:
            lines.append(f"ครูผู้สอนวิชา{first_part}")
        for p in main_parts[1:]:
            lines.append(clean_subject_name(p))
            
    if assist_parts:
        if lines:
            lines[-1] = lines[-1] + " และ"
        first_assist = clean_subject_name(assist_parts[0])
        if first_assist in ["ครูช่วยสอน", "ครูผู้สอน", "ครู"] or not first_assist:
            lines.append("ครูช่วยสอน")
        else:
            lines.append(f"ครูช่วยสอนวิชา{first_assist}")
        for p in assist_parts[1:]:
            lines.append(clean_subject_name(p))
            
    return "<br>".join(lines)

def translate_subject_to_en(subj):
    if not subj:
        return ""
    subj = subj.strip()
    
    # If it contains English characters (like for foreign teachers), extract them
    en_parts = re.findall(r'[a-zA-Z\s&,\(\)/]+', subj)
    if en_parts:
        cleaned = " ".join([p.strip(" ()") for p in en_parts if p.strip(" ()")])
        if cleaned:
            return cleaned
            
    for th_key, en_val in SUBJECT_MAP_EN.items():
        if th_key in subj:
            return en_val
            
    return subj # fallback to original if not mapped

def extract_and_convert_grades(role_th):
    if not role_th:
        return ""
    matches = re.findall(r'([ปม])\.\s*([\d\s\-\,\/]+)', role_th)
    grades = []
    for m in matches:
        level_th = m[0]
        digits = m[1].strip(" ,")
        level_en = 'P.' if level_th == 'ป' else 'M.'
        digits = " ".join(digits.split())
        if digits:
            grades.append(f"{level_en}{digits}")
    return ", ".join(grades) if grades else ""

def format_teaching_role_en(teach, assist):
    main_parts = split_subjects(teach)
    assist_parts = split_subjects(assist)
    
    lines = []
    if main_parts:
        for p in main_parts:
            part_grade = extract_and_convert_grades(p)
            clean_sub = re.sub(r'[ปม]\.\s*[\d\s\-\,\/]+', '', p).strip(" ,")
            sub_en = translate_subject_to_en(clean_sub)
            if sub_en in ["ครูผู้สอน", "Teacher", "ครู"] or not sub_en:
                lines.append("Teacher")
            elif part_grade:
                lines.append(f"{sub_en} Teacher ({part_grade})")
            else:
                lines.append(f"{sub_en} Teacher")
                
    if assist_parts:
        if lines:
            lines[-1] = lines[-1] + " and"
        for p in assist_parts:
            part_grade = extract_and_convert_grades(p)
            clean_sub = re.sub(r'[ปม]\.\s*[\d\s\-\,\/]+', '', p).strip(" ,")
            sub_en = translate_subject_to_en(clean_sub)
            if sub_en in ["ครูช่วยสอน", "ครูผู้สอน", "Teacher", "Teacher Assistant", "ครู"] or not sub_en:
                lines.append("Teacher Assistant")
            elif part_grade:
                lines.append(f"Teacher Assistant ({sub_en} {part_grade})")
            else:
                lines.append(f"Teacher Assistant ({sub_en})")
                
    return "<br>".join(lines) if lines else "Teacher"

print("=== EP Staff Auto Sync ===")

# Parse sync command line arguments or ask user
do_sync = True
if '--sync' in sys.argv:
    do_sync = True
elif '--no-sync' in sys.argv:
    do_sync = False
else:
    if os.path.exists(db_path):
        try:
            ans = input("ต้องการดึงข้อมูลรายชื่อครูจากเซิร์ฟเวอร์โรงเรียน (cpw.ac.th) และ Google Sheet หรือไม่? (y/N): ").strip().lower()
            do_sync = ans in ['y', 'yes', 'ใช่']
        except Exception:
            do_sync = False
    else:
        print("ไม่พบ db_staff.json ระบบจะบังคับดึงข้อมูล (sync) ครั้งแรกอัตโนมัติ...")
        do_sync = True

# 1. Load credentials
fetched_staff = []

if do_sync:
    username = None
    password = None
    if os.path.exists(env_path):
        with open(env_path, 'r', encoding='utf-8') as f:
            for line in f:
                if line.startswith('SCHOOL_USERNAME='):
                    username = line.strip().split('=')[1]
                elif line.startswith('SCHOOL_PASSWORD='):
                    password = line.strip().split('=')[1]
    
    if not username or not password:
        print("Error: Could not find school credentials in EP_StudentSync/.env")
        sys.exit(1)
    
    # 2. Login to school backend
    cj = http.cookiejar.CookieJar()
    opener = urllib.request.build_opener(urllib.request.HTTPCookieProcessor(cj))
    
    login_url = "https://app.cpw.ac.th/cpw_backend/process/login.php"
    login_data = urllib.parse.urlencode({'user': username, 'pass': password}).encode('utf-8')
    
    print("Logging in to cpw.ac.th...")
    req = urllib.request.Request(
        login_url,
        data=login_data,
        headers={
            'Content-Type': 'application/x-www-form-urlencoded',
            'Referer': 'https://app.cpw.ac.th/cpw_backend/login.html',
            'User-Agent': 'Mozilla/5.0'
        }
    )
    
    try:
        with opener.open(req) as resp:
            resp.read()
        print("Login successful.")
    except Exception as e:
        print("Login failed:", e)
        sys.exit(1)
    
    # 3. Fetch staff list
    search_url = "https://app.cpw.ac.th/cpw_backend/controller/php/searcPeople.php"
    search_data = urllib.parse.urlencode({
        'keySearch': '', 'group_id': '', 'rank_id': '', 'de_level': '',
        'de_id': '', 'sub_id': '', 'le_key': '', 'be_id': '',
        'pro_id': '2', # EP Project
        'code': 'searchTime2'
    }).encode('utf-8')
    
    print("Fetching staff list from backend...")
    req2 = urllib.request.Request(
        search_url,
        data=search_data,
        headers={
            'Content-Type': 'application/x-www-form-urlencoded',
            'Referer': 'https://app.cpw.ac.th/cpw_backend/personnel.php',
            'User-Agent': 'Mozilla/5.0'
        }
    )
    
    try:
        with opener.open(req2) as resp:
            html = resp.read().decode('utf-8')
    except Exception as e:
        print("Failed to fetch staff list:", e)
        sys.exit(1)
    
    # 4. Parse HTML rows
    tr_pattern = re.compile(
        r'<tr>.*?alt="(?P<id>\d+)".*?<td>(?P<first>[^<]*)</td>\s*<td>(?P<last>[^<]*)</td>\s*<td>(?P<role>[^<]*)</td>\s*<td>(?P<dept>[^<]*)</td>\s*<td>(?P<subject>[^<]*)</td>',
        re.DOTALL
    )
    
    img_pattern = re.compile(r'img src="([^"]+)"')
    
    for block in re.split(r'</tr>', html):
        match = tr_pattern.search(block)
        if match:
            p_id = match.group('id')
            first = match.group('first').strip()
            last = match.group('last').strip()
            role = match.group('role').strip()
            dept = match.group('dept').strip()
            subj = match.group('subject').strip()
            
            img_match = img_pattern.search(block)
            img_src = img_match.group(1) if img_match else ""
            if img_src.startswith('../'):
                img_url = "https://app.cpw.ac.th/" + img_src[3:]
            else:
                img_url = img_src
                
            fetched_staff.append({
                'id': p_id,
                'first': first,
                'last': last,
                'fullName': f"{first} {last}".strip(),
                'role': role,
                'dept': dept,
                'subj': subj,
                'img_url': img_url
            })
    
    print(f"Parsed {len(fetched_staff)} staff members from backend.")

# 5. Load existing db_staff.json
existing_staff = []
if os.path.exists(db_path):
    try:
        with open(db_path, 'r', encoding='utf-8') as f:
            existing_staff = json.load(f)
        print(f"Loaded {len(existing_staff)} existing staff from db_staff.json.")
    except Exception as e:
        print("Failed to read db_staff.json, starting fresh:", e)

# 5.5 Download and load Google Sheet data (AY2569) to override roles/names
google_staff_map = {}
homeroom_map = {}

if do_sync:
    import urllib.request
    import openpyxl
    import re
    
    gs_url = "https://docs.google.com/spreadsheets/d/1GW5GwqvwRI3Dvz0YcYvEOTy4eemXoewb7JTUlFYvc7o/export?format=xlsx"
    temp_xlsx_path = "temp_sync_sheet.xlsx"
    
    try:
        print("Downloading Google Sheet (AY2569)...")
        req_gs = urllib.request.Request(gs_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req_gs) as resp:
            gs_data = resp.read()
        with open(temp_xlsx_path, 'wb') as f:
            f.write(gs_data)
        print("Download successful.")
        
        wb = openpyxl.load_workbook(temp_xlsx_path, data_only=True)
        
        # 1. Parse general overrides from 'Assignments' sheet FIRST (defines default classes)
        sheet_assign = wb['Assignments'] if 'Assignments' in wb.sheetnames else wb.worksheets[4]
        for row in sheet_assign.iter_rows(values_only=True):
            if not row or len(row) < 11:
                continue
            p_id_val = str(row[1]).split('.')[0].strip() if row[1] is not None else ""
            if p_id_val and p_id_val.isdigit():
                name_th_val = str(row[2]).strip() if row[2] is not None else ""
                pos_val = str(row[3]).strip() if row[3] is not None else ""
                name_en_val = str(row[4]).strip() if row[4] is not None else ""
                nick_val = str(row[5]).strip() if row[5] is not None else ""
                admin_val = str(row[6]).strip() if row[6] is not None else ""
                teach_val = str(row[7]).strip() if row[7] is not None else ""
                assist_val = str(row[8]).strip() if row[8] is not None else ""
                level_val = str(row[9]).strip() if row[9] is not None else ""
                room_val = str(row[10]).strip() if row[10] is not None else ""
                
                google_staff_map[p_id_val] = {
                    'name_th': name_th_val,
                    'name_en': name_en_val,
                    'nickname': nick_val,
                    'admin_role': admin_val,
                    'teaching_role': teach_val,
                    'assist_teaching': assist_val
                }
                
                if level_val and room_val:
                    room_clean = room_val.split('.')[0].strip()
                    classroom_str = f"{level_val}/{room_clean}"
                    homeroom_map[p_id_val] = classroom_str
                    
        print(f"Loaded {len(google_staff_map)} staff metadata records from Google Sheet 'Assignments'.")
        
        # 2. Parse homeroom mapping from 'For Mameow' sheet SECOND (overrides defaults with actual transfers)
        sheet_hr = wb['For Mameow'] if 'For Mameow' in wb.sheetnames else wb.worksheets[0]
        for row in sheet_hr.iter_rows(values_only=True):
            if not row or len(row) < 4:
                continue
            level_val = str(row[0]).strip() if row[0] is not None else ""
            room_val = str(row[1]).strip() if row[1] is not None else ""
            p_id_val = str(row[2]).strip() if row[2] is not None else ""
            
            p_id_val = p_id_val.split('.')[0].strip()
            if p_id_val and p_id_val.isdigit() and level_val and room_val:
                room_clean = room_val.split('.')[0].strip()
                classroom_str = f"{level_val}/{room_clean}"
                homeroom_map[p_id_val] = classroom_str
    
        print(f"Loaded {len(homeroom_map)} homeroom mappings from Google Sheet 'For Mameow'.")
        
        if os.path.exists(temp_xlsx_path):
            os.remove(temp_xlsx_path)
            
    except Exception as e:
        print(f"Failed to fetch or parse Google Sheet: {e}")
        if os.path.exists(temp_xlsx_path):
            os.remove(temp_xlsx_path)

# Helper function to find existing entry
def find_existing(fullName, p_id):
    for x in existing_staff:
        img_name = os.path.basename(x.get('imageUrl', ''))
        if img_name.startswith(p_id):
            return x
        if x.get('name') == fullName:
            return x
    return None

# 5.7 Run portrait processor to update transparent PNGs from raw images (always run)
print("Running portrait processor (process_portraits.py)...")
try:
    subprocess.run([sys.executable, 'process_portraits.py'], check=True)
except Exception as e:
    print("Portrait processing failed:", e)

# 6. Merge & Update
updated_staff = []
os.makedirs(img_dir, exist_ok=True)

if do_sync:
    print("ระบบรันโหมด Sync: กำลังจัดเตรียมและซิงค์ข้อมูลบุคคล...")
    for num, fs in enumerate(fetched_staff):
        fullName = fs['fullName']
        p_id = fs['id']
        
        is_foreign = bool(re.search(r'[a-zA-Z]', fullName))
        gs_info = google_staff_map.get(p_id, {})
        
        # 1. Resolve Name
        if gs_info.get('name_th'):
            fullName = gs_info['name_th']
        name_en = gs_info.get('name_en', '')
        if not name_en and is_foreign:
            name_en = fullName
            
        # 2. Resolve Role & Category
        role = fs['role']
        category = ""
        
        admin_role_th = gs_info.get('admin_role', '')
        is_admin = bool(admin_role_th) or p_id in ['10997', '10158', '10092'] or p_id.startswith('2')
        
        if is_admin:
            category = 'admin'
            if admin_role_th:
                role = admin_role_th
            else:
                if p_id == '10997': role = 'ที่ปรึกษาโครงการฯ'
                elif p_id == '10158': role = 'หัวหน้าโครงการฯ'
                elif p_id == '10092': role = 'ผู้ช่วยหัวหน้าโครงการฯ'
                elif p_id == '10179': role = 'หัวหน้างานปกครอง โครงการฯ'
                else: role = 'เจ้าหน้าที่'
        else:
            if is_foreign:
                category = 'foreign'
            elif 'มัธยม' in fs['dept'] or 'ม.' in fs['dept']:
                category = 'secondary'
            else:
                category = 'primary'
                
        # Define normalized schema fields
        role_subject = ""
        role_subject_en = ""
        role_homeroom = ""
        role_homeroom_en = ""
        role_en = 'Teacher'
        classRoom_th = fs['dept']
        classRoom_en = 'EP Classroom'
        if 'ประถม' in classRoom_th:
            classRoom_en = 'Primary Level'
        elif 'มัธยม' in classRoom_th or 'ม.' in classRoom_th:
            classRoom_en = 'Secondary Level'
        
        if category == 'admin':
            role_en = ADMIN_ROLE_EN_MAP.get(role, 'Support Staff')
        else:
            teach_val = gs_info.get('teaching_role', '')
            assist_val = gs_info.get('assist_teaching', '')
            
            if teach_val or assist_val:
                role_subject = format_teaching_role(teach_val, assist_val)
                role_subject_en = format_teaching_role_en(teach_val, assist_val)
            else:
                role_subject = f"ครูผู้สอน"
                role_subject_en = "Subject Teacher"
                
            role_en = role_subject_en
            role = role_subject
            
        if p_id in homeroom_map:
            classroom_str = homeroom_map[p_id]
            role_homeroom = f"ครูประจำชั้น {classroom_str}"
            role_homeroom_en = f"Advisor {classroom_str.replace('ป.', 'P.').replace('ม.', 'M.')}"
            classRoom_th = classroom_str
            classRoom_en = classroom_str.replace('ป.', 'P.').replace('ม.', 'M.')

        # Resolve Image (Only search for processed -2.png or raw -2.jpg/jpeg)
        img_filename = "default.jpg" # default fallback
        for ext in ['-2.png', '-2.jpg', '-2.jpeg']:
            candidate = f"{p_id}{ext}"
            if os.path.exists(os.path.join(img_dir, candidate)):
                img_filename = candidate
                break
            elif os.path.exists(os.path.join(img_dir, 'raw', candidate)):
                img_filename = f"raw/{candidate}"
                break
            
        web_img_path = f"img/teachers/{img_filename}"

        match = find_existing(fullName, p_id)
        if match:
            match['name'] = fullName
            match['name_en'] = name_en
            match['personnelId'] = p_id
            match['role'] = role
            match['role_en'] = role_en
            match['role_subject'] = role_subject
            match['role_subject_en'] = role_subject_en
            match['role_homeroom'] = role_homeroom
            match['role_homeroom_en'] = role_homeroom_en
            match['classRoom'] = classRoom_th
            match['classRoom_en'] = classRoom_en
            match['category'] = category
            match['imageUrl'] = web_img_path
            updated_staff.append(match)
        else:
            new_record = {
                'name': fullName,
                'name_en': name_en,
                'personnelId': p_id,
                'role': role,
                'role_en': role_en,
                'role_subject': role_subject,
                'role_subject_en': role_subject_en,
                'role_homeroom': role_homeroom,
                'role_homeroom_en': role_homeroom_en,
                'category': category,
                'isForeign': is_foreign,
                'classRoom': classRoom_th,
                'classRoom_en': classRoom_en,
                'subjects': [fs['subj']] if fs['subj'] else [],
                'subjects_en': [],
                'imageUrl': web_img_path,
                'sortOrder': len(existing_staff) + num + 1
            }
            print(f"Adding new staff member: {fullName} ({category})")
            updated_staff.append(new_record)
else:
    print("ระบบรันโหมด Local: ใช้รายชื่อเดิมจากฐานข้อมูลท้องถิ่น และเชื่อมโยงภาพถ่ายใหม่...")
    updated_staff = existing_staff
    for member in updated_staff:
        p_id = member.get('personnelId', '')
        if not p_id:
            continue
            
        img_filename = "default.jpg" # default fallback
        for ext in ['-2.png', '-2.jpg', '-2.jpeg']:
            candidate = f"{p_id}{ext}"
            if os.path.exists(os.path.join(img_dir, candidate)):
                img_filename = candidate
                break
            elif os.path.exists(os.path.join(img_dir, 'raw', candidate)):
                img_filename = f"raw/{candidate}"
                break
                
        member['imageUrl'] = f"img/teachers/{img_filename}"
        
        # Translate roles locally to ensure clean English display
        is_foreign = member.get('isForeign', False) or bool(re.search(r'[a-zA-Z]', member.get('name', '')))
        category = member.get('category', '')
        
        if category == 'admin':
            member['role_en'] = ADMIN_ROLE_EN_MAP.get(member.get('role', ''), 'Support Staff')
        else:
            role_th = member.get('role', '')
            if role_th:
                clean_role = role_th.replace('<br>', ',').replace('\n', ',')
                parts = [p.strip() for p in clean_role.split(',') if p.strip()]
                
                teach_subs = []
                assist_subs = []
                for p in parts:
                    p_clean = p.replace(' และ', '').strip()
                    if 'ครูผู้สอนวิชา' in p_clean:
                        teach_subs.append(p_clean.replace('ครูผู้สอนวิชา', '').strip())
                    elif 'ครูช่วยสอนวิชา' in p_clean:
                        assist_subs.append(p_clean.replace('ครูช่วยสอนวิชา', '').strip())
                    else:
                        if assist_subs:
                            assist_subs.append(p_clean)
                        elif teach_subs:
                            teach_subs.append(p_clean)
                        else:
                            teach_subs.append(p_clean)
                            
                teach_val = ", ".join(teach_subs)
                assist_val = ", ".join(assist_subs)
                
                if teach_val or assist_val:
                    role_en = format_teaching_role_en(teach_val, assist_val)
                    role_th_new = format_teaching_role(teach_val, assist_val)
                else:
                    role_en = 'Subject Teacher'
                    role_th_new = role_th
                
                member['role'] = role_th_new
                member['role_subject'] = role_th_new
                member['role_subject_en'] = role_en
                member['role_en'] = role_en

# 6.5 Sort staff according to handbook hierarchy
def get_sort_key(x):
    p_id = x.get('personnelId', '')
    if not p_id:
        img_name = os.path.basename(x.get('imageUrl', ''))
        p_id = "".join(filter(str.isdigit, img_name))
        
    is_foreign = x.get('isForeign', False)
    category = x.get('category', '')
    role = x.get('role', '')
    
    # 1. Top Administrators (Advisor, Head)
    if p_id == '10997':
        return (0, 0, 0, 0)
    elif p_id == '10158':
        return (0, 1, 0, 0)
        
    # 2. Other admin/office support
    if category == 'admin':
        admin_order = ['10092', '10081', '10179', '20002', '20008', '20040', '20031', '20038', '20033', '20009', '20023', '20041']
        if p_id in admin_order:
            return (1, admin_order.index(p_id), 0, 0)
        else:
            return (1, 100 + int(p_id or 9999), 0, 0)
            
    # Check for homeroom teachers sorting (Group 2)
    # E.g. 'ครูประจำชั้น ป.1/14' -> (2, level, grade, room)
    if 'ครูประจำชั้น' in role:
        m = re.search(r'([ปม])\.(\d+)/(\d+)', role)
        if m:
            level = 0 if m.group(1) == 'ป' else 1
            grade = int(m.group(2))
            room = int(m.group(3))
            return (2, level, grade, room)
            
    # 3. Thai Subject Teachers (who are not homeroom)
    if not is_foreign:
        return (3, int(p_id or 99999), 0, 0)
        
    # 4. Foreign Subject Teachers
    return (4, int(p_id or 99999), 0, 0)

print("Sorting staff according to handbook hierarchy...")
updated_staff.sort(key=get_sort_key)

# Re-assign sortOrder based on the sorted list and inject educational profile data for details modal
for idx, x in enumerate(updated_staff):
    x['sortOrder'] = idx + 1
    
    p_id = x.get('personnelId', '')
    is_foreign = x.get('isForeign', False)
    
    if p_id == '10997':
        x['education'] = "ครุศาสตรบัณฑิต (ค.บ.) สาขาวิชาการสอนภาษาอังกฤษ, จุฬาลงกรณ์มหาวิทยาลัย"
        x['education_en'] = "B.Ed. in English Language Teaching, Chulalongkorn University"
        x['experience'] = "ประสบการณ์การสอนและบริหารหลักสูตรภาษาอังกฤษมากกว่า 25 ปี"
        x['experience_en'] = "Over 25 years of teaching and managing English Program curricula"
        x['nationality_en'] = "Thai"
    elif p_id == '10158':
        x['education'] = "ศึกษาศาสตรมหาบัณฑิต (ศษ.ม.) สาขาการบริหารการศึกษา, มหาวิทยาลัยเกษตรศาสตร์"
        x['education_en'] = "M.Ed. in Educational Administration, Kasetsart University"
        x['experience'] = "ประสบการณ์การบริหารงานโครงการส่งเสริมความเป็นเลิศทางภาษาอังกฤษ 15 ปี"
        x['experience_en'] = "15 years of experience in administering English Program excellence tracks"
        x['nationality_en'] = "Thai"
    elif p_id == '10092':
        x['education'] = "วิทยาศาสตรบัณฑิต (วท.บ.) สาขาวิชาคณิตศาสตร์, มหาวิทยาลัยธรรมศาสตร์"
        x['education_en'] = "B.Sc. in Mathematics, Thammasat University"
        x['experience'] = "ครูผู้สอนวิชาคณิตศาสตร์ระดับมัธยมศึกษาตอนปลาย ประสบการณ์สอน 10 ปี"
        x['experience_en'] = "Senior high school Mathematics teacher with 10 years of teaching experience"
        x['nationality_en'] = "Thai"
    else:
        if 'education' not in x or not x['education']:
            if is_foreign:
                x['education'] = "ปริญญาตรีในสาขาวิชาเฉพาะทาง และใบอนุญาตประกอบวิชาชีพครูสากล"
                x['education_en'] = "Bachelor's Degree in specialized field & international teaching credentials"
                x['experience'] = "มีประสบการณ์สอนหลักสูตรนานาชาติ หรือหลักสูตรสองภาษา (Bilingual/EP)"
                x['experience_en'] = "Experienced in teaching international or bilingual (EP) programs"
                x['nationality_en'] = "Native English Speaker / International Faculty"
            else:
                x['education'] = "ปริญญาตรีทางการศึกษา (ศษ.บ. / ค.บ.) หรือสาขาวิชาเฉพาะทาง"
                x['education_en'] = "Bachelor's Degree in Education (B.Ed.) or specialized field"
                x['experience'] = "ผ่านการประเมินมาตรฐานวิชาชีพครูและการสอนสองภาษาของโครงการ"
                x['experience_en'] = "Certified under English Program teaching standards and professional assessment"
                x['nationality_en'] = "Thai"

# Save db_staff.json
with open(db_path, 'w', encoding='utf-8') as f:
    json.dump(updated_staff, f, ensure_ascii=False, indent=2)

print(f"Successfully synced and sorted {len(updated_staff)} records to db_staff.json.")

# 7. Run compilation
print("Running static compiler (compile.js)...")
try:
    subprocess.run(['node', 'compile.js'], check=True)
    print("Static build compiled successfully!")
except Exception as e:
    print("Compilation failed:", e)
