import os
import sys
import json
import urllib.request
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

# Paths
base_dir = os.path.dirname(os.path.abspath(__file__))
settings_path = os.path.join(base_dir, 'data', 'db_site_settings.json')
slides_path = os.path.join(base_dir, 'data', 'db_slides.json')

gs_url = "https://docs.google.com/spreadsheets/d/1GW5GwqvwRI3Dvz0YcYvEOTy4eemXoewb7JTUlFYvc7o/export?format=xlsx"
temp_xlsx_path = "temp_sync_settings_sheet.xlsx"

print("=== EP Site Settings & Slides Sync ===")

try:
    print("Downloading Google Sheet...")
    req_gs = urllib.request.Request(gs_url, headers={'User-Agent': 'Mozilla/5.0'})
    with urllib.request.urlopen(req_gs) as resp:
        gs_data = resp.read()
    with open(temp_xlsx_path, 'wb') as f:
        f.write(gs_data)
    print("Download successful.")
    
    wb = openpyxl.load_workbook(temp_xlsx_path, data_only=True)
    
    # Load existing settings
    existing_settings = {}
    if os.path.exists(settings_path):
        with open(settings_path, 'r', encoding='utf-8') as f:
            existing_settings = json.load(f)
            
    # Load existing slides
    existing_slides = []
    if os.path.exists(slides_path):
        with open(slides_path, 'r', encoding='utf-8') as f:
            existing_slides = json.load(f)

    # 1. Parse 'SiteSettings' key-value sheet
    if 'SiteSettings' in wb.sheetnames:
        sheet = wb['SiteSettings']
        print("Parsing 'SiteSettings'...")
        
        # Ensure structures exist
        if 'about' not in existing_settings:
            existing_settings['about'] = {}
        if 'fees' not in existing_settings:
            existing_settings['fees'] = {
                "primary": {"tuition": "0", "food": "0", "astro": "0", "activities": "0"},
                "junior": {"tuition": "0", "food": "0", "astro": "0", "activities": "0"},
                "senior": {"tuition": "0", "food": "0", "astro": "0", "activities": "0"}
            }
        if 'contact' not in existing_settings:
            existing_settings['contact'] = {}

        for row in sheet.iter_rows(values_only=True):
            if not row or len(row) < 2:
                continue
            key = str(row[0]).strip() if row[0] is not None else ""
            val = str(row[1]).strip() if row[1] is not None else ""
            
            if not key:
                continue
                
            # Map key-value to json structure
            if key == 'about_intro_th':
                existing_settings['about']['intro_th'] = val
            elif key == 'about_intro_en':
                existing_settings['about']['intro_en'] = val
            elif key == 'contact_phone':
                existing_settings['contact']['phone'] = val
            elif key == 'contact_email':
                existing_settings['contact']['email'] = val
            elif key == 'contact_facebook':
                existing_settings['contact']['facebook'] = val
            elif key == 'contact_address_th':
                existing_settings['contact']['address_th'] = val
            elif key == 'contact_address_en':
                existing_settings['contact']['address_en'] = val
                
            # Tuition summary cards
            elif key.startswith('fees_'):
                parts = key.split('_')  # e.g., fees_primary_tuition
                if len(parts) == 3:
                    lvl = parts[1]  # primary, junior, senior
                    field = parts[2]  # tuition, food, astro, activities
                    if lvl in existing_settings['fees'] and field in existing_settings['fees'][lvl]:
                        existing_settings['fees'][lvl][field] = val

    # 2. Parse 'TuitionFees' sheet
    if 'TuitionFees' in wb.sheetnames:
        sheet = wb['TuitionFees']
        print("Parsing 'TuitionFees'...")
        fees_table = []
        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if idx == 0:  # skip header row
                continue
            if not row or len(row) < 6 or not row[0]:
                continue
            
            level_th = str(row[0]).strip()
            level_en = str(row[1]).strip() if row[1] is not None else ""
            tuition = str(row[2]).strip() if row[2] is not None else "0"
            food = str(row[3]).strip() if row[3] is not None else "-"
            astro = str(row[4]).strip() if row[4] is not None else "0"
            total = str(row[5]).strip() if row[5] is not None else "0"
            
            fees_table.append({
                "level_th": level_th,
                "level_en": level_en,
                "tuition": tuition,
                "food": food,
                "astro": astro,
                "total": total
            })
        if fees_table:
            existing_settings['fees_table'] = fees_table

    # 3. Parse 'BookPrices' sheet
    if 'BookPrices' in wb.sheetnames:
        sheet = wb['BookPrices']
        print("Parsing 'BookPrices'...")
        books_table = []
        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if idx == 0:  # skip header row
                continue
            if not row or len(row) < 4 or not row[0]:
                continue
            
            level_th = str(row[0]).strip()
            level_en = str(row[1]).strip() if row[1] is not None else ""
            plan = str(row[2]).strip() if row[2] is not None else "-"
            price = str(row[3]).strip() if row[3] is not None else "0"
            
            books_table.append({
                "level_th": level_th,
                "level_en": level_en,
                "plan": plan,
                "price": price
            })
        if books_table:
            existing_settings['books_table'] = books_table

    # 4. Parse 'History' timeline sheet
    if 'History' in wb.sheetnames:
        sheet = wb['History']
        print("Parsing 'History'...")
        history = []
        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if idx == 0:  # skip header row
                continue
            if not row or len(row) < 3 or not row[0]:
                continue
            
            year = str(row[0]).strip()
            title = str(row[1]).strip()
            desc = str(row[2]).strip() if row[2] is not None else ""
            
            history.append({
                "year": year,
                "title": title,
                "desc": desc
            })
        if history:
            if 'about' not in existing_settings:
                existing_settings['about'] = {}
            existing_settings['about']['history'] = history

    # 5. Parse 'Slides' sheet
    if 'Slides' in wb.sheetnames:
        sheet = wb['Slides']
        print("Parsing 'Slides'...")
        slides = []
        for idx, row in enumerate(sheet.iter_rows(values_only=True)):
            if idx == 0:  # skip header row
                continue
            if not row or len(row) < 13 or not row[2]:  # title is required
                continue
            
            badge = str(row[0]).strip() if row[0] is not None else ""
            badge_en = str(row[1]).strip() if row[1] is not None else ""
            title = str(row[2]).strip()
            title_en = str(row[3]).strip() if row[3] is not None else title
            desc = str(row[4]).strip() if row[4] is not None else ""
            desc_en = str(row[5]).strip() if row[5] is not None else desc
            
            btn1Text = str(row[6]).strip() if row[6] is not None else ""
            btn1TextEn = str(row[7]).strip() if row[7] is not None else ""
            btn1Link = str(row[8]).strip() if row[8] is not None else ""
            
            btn2Text = str(row[9]).strip() if row[9] is not None else ""
            btn2TextEn = str(row[10]).strip() if row[10] is not None else ""
            btn2Link = str(row[11]).strip() if row[11] is not None else ""
            
            img_url = str(row[12]).strip() if row[12] is not None else ""
            
            slides.append({
                "badge": badge,
                "badge_en": badge_en,
                "title": title,
                "title_en": title_en,
                "description": desc,
                "description_en": desc_en,
                "btn1Text": btn1Text,
                "btn1TextEn": btn1TextEn,
                "btn1Link": btn1Link,
                "btn2Text": btn2Text,
                "btn2TextEn": btn2TextEn,
                "btn2Link": btn2Link,
                "imageUrl": img_url
            })
        if slides:
            existing_slides = slides

    # Save outputs
    os.makedirs(os.path.dirname(settings_path), exist_ok=True)
    
    with open(settings_path, 'w', encoding='utf-8') as f:
        json.dump(existing_settings, f, ensure_ascii=False, indent=2)
    print(f"Saved site settings to: {settings_path}")

    with open(slides_path, 'w', encoding='utf-8') as f:
        json.dump(existing_slides, f, ensure_ascii=False, indent=2)
    print(f"Saved slides banner to: {slides_path}")

    # Copy to preview directory if it exists
    preview_dir = os.path.join(base_dir, 'preview', 'data')
    if os.path.exists(os.path.dirname(preview_dir)):
        os.makedirs(preview_dir, exist_ok=True)
        with open(os.path.join(preview_dir, 'db_site_settings.json'), 'w', encoding='utf-8') as f:
            json.dump(existing_settings, f, ensure_ascii=False, indent=2)
        with open(os.path.join(preview_dir, 'db_slides.json'), 'w', encoding='utf-8') as f:
            json.dump(existing_slides, f, ensure_ascii=False, indent=2)
        print("Copied database files to preview/data/")

    if os.path.exists(temp_xlsx_path):
        os.remove(temp_xlsx_path)

except Exception as e:
    print(f"Error syncing site settings/slides: {e}")
    if os.path.exists(temp_xlsx_path):
        os.remove(temp_xlsx_path)
    sys.exit(1)
