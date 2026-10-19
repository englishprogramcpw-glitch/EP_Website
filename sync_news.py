import os
import re
import ssl
import json
import urllib.request

def fetch_html(url):
    print(f"Fetching: {url}")
    try:
        # Create unverified context to bypass any SSL certificate issues
        context = ssl._create_unverified_context()
        req = urllib.request.Request(
            url, 
            headers={'User-Agent': 'Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/100.0.0.0 Safari/537.36'}
        )
        with urllib.request.urlopen(req, context=context, timeout=10) as response:
            return response.read().decode('utf-8', errors='ignore')
    except Exception as e:
        print(f"Error fetching {url}: {e}")
        return ""

def sync_news():
    # Keywords to filter EP-related content
    ep_keywords = ["EP", "English Program", "ภาคภาษาอังกฤษ", "อีพี", "ห้องเรียนพิเศษภาษาอังกฤษ"]
    
    # 1. Scrape News (Announcements)
    news_html = fetch_html("https://www.cpw.ac.th/news_public.php")
    announcements = []
    
    if news_html:
        # Match news items: class="new2", title="...", href="..."
        pattern = re.compile(
            r'<div class="new2" title="([^"]+)">\s*<a href="([^"]+)"[^>]*>\s*&nbsp;<span class="newsname2">.*?</span>\s*<span class="newdate2">([^<]+)</span>',
            re.DOTALL
        )
        matches = pattern.findall(news_html)
        print(f"Found {len(matches)} total public announcements on main site.")
        
        for title, url, date in matches:
            title = title.strip()
            date = date.strip()
            url = url.strip()
            
            # Check if title matches EP keywords
            is_ep = any(kw.lower() in title.lower() for kw in ep_keywords)
            if is_ep:
                # Classify category based on keywords
                category = "กิจกรรม"
                title_lower = title.lower()
                
                # Admissions keywords
                if any(k in title_lower for k in ["รับสมัคร", "สมัคร", "เข้าเรียน", "สิทธิ์เข้าเรียน", "admission", "apply"]):
                    category = "รับสมัครเรียน"
                # Awards keywords
                elif any(k in title_lower for k in ["ผลงาน", "คะแนน", "รางวัล", "ชนะเลิศ", "ยินดี", "o-net", "award", "contest", "congratulations"]):
                    category = "ผลงานนักเรียน"
                    
                description = f"ประกาศอย่างเป็นทางการเกี่ยวกับโครงการภาคภาษาอังกฤษ (EP) เรื่อง: {title} สามารถคลิกปุ่มด้านล่างเพื่อเปิดอ่านประกาศฉบับเต็มในรูปแบบ PDF ได้เลยค่ะ"
                
                announcements.append({
                    "title": title,
                    "title_en": title,
                    "category": category,
                    "date": date,
                    "url": url,
                    "content": description,
                    "content_en": description,
                    "imageUrl": "" # Text PDF announcements do not have custom images
                })
                
    # 2. Scrape Activities (Albums across multiple pages)
    activities = []
    seen_urls = set()
    pattern = re.compile(
        r'<div class="Gallerygroups"><a href="([^"]+)"><div class="Galleryboximage" title="([^"]+)"><img src="([^"]*)"[^>]*/>\s*</div><div class="Galleryname">.*?</div><div class="Gallerydate">\s*วันที่\s*([^<]+)</div>\s*<div class="Gallerypicnumber"><i class="fa fa-picture-o" aria-hidden="true"></i>\s*(\d+)\s*รูป</div>',
        re.DOTALL
    )

    # Scan pages 1 to 15 to collect comprehensive activity albums
    for page in range(1, 16):
        page_url = f"https://www.cpw.ac.th/albums.php?page={page}" if page > 1 else "https://www.cpw.ac.th/albums.php"
        albums_html = fetch_html(page_url)
        if not albums_html:
            break
        
        matches = pattern.findall(albums_html)
        if not matches:
            break
            
        print(f"Page {page}: found {len(matches)} total albums.")
        ep_in_page = 0
        
        for rel_url, title, img_url, date, photo_count in matches:
            full_url = f"https://www.cpw.ac.th/{rel_url.strip()}"
            if full_url in seen_urls:
                continue
            seen_urls.add(full_url)
            
            title = title.replace('###s', "'s").strip()
            img_url = img_url.strip()
            date = date.strip()
            photo_count = int(photo_count)
            
            # Check if title matches EP keywords
            is_ep = any(kw.lower() in title.lower() for kw in ep_keywords)
            if is_ep:
                ep_in_page += 1
                # Scrape all image URLs inside this specific album
                images = []
                sub_html = fetch_html(full_url)
                if sub_html:
                    img_pattern = re.compile(r'<figure><a href="([^"]+)" itemprop="contentUrl"')
                    images = img_pattern.findall(sub_html)
                    # Convert to full URL if relative
                    images = [img if img.startswith('http') else f"https://www.cpw.ac.th/{img}" for img in images]
                    try:
                        print(f"  [+] Page {page}: {len(images)} images -> {title}")
                    except Exception:
                        print(f"  [+] Page {page}: {len(images)} images -> id: {full_url.split('=')[-1]}")
                
                activities.append({
                    "title": title,
                    "date": date,
                    "url": full_url,
                    "imageUrl": img_url if img_url else "https://images.unsplash.com/photo-1511632765486-a01980e01a18?w=800&auto=format&fit=crop&q=60",
                    "photoCount": photo_count,
                    "images": images
                })
        print(f"Page {page}: collected {ep_in_page} EP albums. Total so far: {len(activities)}")

    # 3. Read custom overrides and announcements from Google Sheet (if tabs exist)
    gs_url = "https://docs.google.com/spreadsheets/d/1GW5GwqvwRI3Dvz0YcYvEOTy4eemXoewb7JTUlFYvc7o/export?format=xlsx"
    temp_xlsx_path = "temp_sync_news_sheet.xlsx"
    
    try:
        print("Downloading Google Sheet for custom news/activities overrides...")
        req_gs = urllib.request.Request(gs_url, headers={'User-Agent': 'Mozilla/5.0'})
        with urllib.request.urlopen(req_gs) as resp:
            gs_data = resp.read()
        with open(temp_xlsx_path, 'wb') as f:
            f.write(gs_data)
        
        import openpyxl
        wb = openpyxl.load_workbook(temp_xlsx_path, data_only=True)
        
        # Parse 'News' tab
        if 'News' in wb.sheetnames:
            sheet_news = wb['News']
            custom_count = 0
            for idx, row in enumerate(sheet_news.iter_rows(values_only=True)):
                if idx == 0:  # skip header row
                    continue
                if not row or not row[2]:  # title is required
                    continue
                
                date_val = str(row[0]).strip() if row[0] is not None else ""
                cat_val = str(row[1]).strip() if row[1] is not None else "กิจกรรม"
                title_val = str(row[2]).strip() if row[2] is not None else ""
                title_en_val = str(row[3]).strip() if row[3] is not None else title_val
                content_val = str(row[4]).strip() if row[4] is not None else ""
                content_en_val = str(row[5]).strip() if row[5] is not None else content_val
                url_val = str(row[6]).strip() if row[6] is not None else ""
                img_url_val = str(row[7]).strip() if row[7] is not None else ""
                
                # Unshift so custom news appears at the top
                announcements.insert(0, {
                    "title": title_val,
                    "title_en": title_en_val,
                    "category": cat_val,
                    "date": date_val,
                    "url": url_val,
                    "content": content_val,
                    "content_en": content_en_val,
                    "imageUrl": img_url_val
                })
                custom_count += 1
            print(f"Loaded {custom_count} custom announcements from Google Sheet 'News' tab.")

        # Parse 'Activities' tab
        if 'Activities' in wb.sheetnames:
            sheet_act = wb['Activities']
            custom_act_count = 0
            for idx, row in enumerate(sheet_act.iter_rows(values_only=True)):
                if idx == 0:  # skip header row
                    continue
                if not row or not row[1]:  # title is required
                    continue
                
                date_val = str(row[0]).strip() if row[0] is not None else ""
                title_val = str(row[1]).strip() if row[1] is not None else ""
                cover_img_val = str(row[2]).strip() if row[2] is not None else ""
                photo_count_val = int(row[3]) if row[3] is not None and str(row[3]).isdigit() else 0
                img_list_raw = str(row[4]).strip() if row[4] is not None else ""
                
                images = [url.strip() for url in re.split(r'[,，\n]', img_list_raw) if url.strip()]
                if not photo_count_val:
                    photo_count_val = len(images) if images else 1
                
                activities.insert(0, {
                    "title": title_val,
                    "date": date_val,
                    "url": "news.php?tab=activities",
                    "imageUrl": cover_img_val or (images[0] if images else "https://images.unsplash.com/photo-1511632765486-a01980e01a18?w=800&auto=format&fit=crop&q=60"),
                    "photoCount": photo_count_val,
                    "images": images
                })
                custom_act_count += 1
            print(f"Loaded {custom_act_count} custom activities from Google Sheet 'Activities' tab.")
            
        if os.path.exists(temp_xlsx_path):
            os.remove(temp_xlsx_path)
            
    except Exception as e:
        print(f"Notice: Google Sheet sync skipped or failed (will use public site scraper only): {e}")
        if os.path.exists(temp_xlsx_path):
            os.remove(temp_xlsx_path)

    # Save to data directory
    base_dir = os.path.dirname(os.path.abspath(__file__))
    data_dir = os.path.join(base_dir, 'data')
    os.makedirs(data_dir, exist_ok=True)
    
    # Save announcements
    news_file = os.path.join(data_dir, 'db_news.json')
    with open(news_file, 'w', encoding='utf-8') as f:
        json.dump(announcements, f, ensure_ascii=False, indent=2)
    print(f"Saved {len(announcements)} EP announcements to: {news_file}")
    
    # Save activities
    activities_file = os.path.join(data_dir, 'db_activities.json')
    with open(activities_file, 'w', encoding='utf-8') as f:
        json.dump(activities, f, ensure_ascii=False, indent=2)
    print(f"Saved {len(activities)} EP activities to: {activities_file}")
    
    # Copy to preview directory if it exists
    preview_data_dir = os.path.join(base_dir, 'preview', 'data')
    if os.path.exists(os.path.dirname(preview_data_dir)):
        os.makedirs(preview_data_dir, exist_ok=True)
        with open(os.path.join(preview_data_dir, 'db_news.json'), 'w', encoding='utf-8') as f:
            json.dump(announcements, f, ensure_ascii=False, indent=2)
        with open(os.path.join(preview_data_dir, 'db_activities.json'), 'w', encoding='utf-8') as f:
            json.dump(activities, f, ensure_ascii=False, indent=2)
        print("Successfully synced JSON databases to preview/data/")

if __name__ == '__main__':
    sync_news()
